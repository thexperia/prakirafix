"""Pemeriksa file Excel upload. Mengembalikan error (menghalangi), peringatan (tidak menghalangi), dan data bersih."""
import io
import re

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

MIN_OBS = 24
Q_OK = {"Q1", "Q2", "Q3", "Q4"}


def pilih_sheet(xls):
    names = xls.sheet_names
    if "Data" in names:
        return "Data"
    for n in names:
        if n.lower() not in ("petunjuk", "contoh", "metadata"):
            return n
    return names[0]


def periksa(file_bytes):
    """Hasil: dict(ok, sheet, errors, warnings, df, labels, info)."""
    errors, warnings = [], []
    res = {"ok": False, "sheet": None, "errors": errors, "warnings": warnings, "df": None, "labels": {}, "info": {}}
    try:
        xls = pd.ExcelFile(io.BytesIO(file_bytes))
    except Exception:
        errors.append({"lokasi": "File", "pesan": "File tidak bisa dibaca sebagai Excel (.xlsx). Pastikan memakai template."})
        return res
    sheet = pilih_sheet(xls)
    res["sheet"] = sheet
    raw = pd.read_excel(xls, sheet_name=sheet, header=None)
    if raw.shape[0] < 4 or raw.shape[1] < 3:
        errors.append({"lokasi": f"Sheet {sheet}", "pesan": "Sheet terlalu kecil. Baris 1 nama indikator, baris 2 kode, data mulai baris 3, kolom A tahun, kolom B triwulan, indikator mulai kolom C."})
        return res

    # --- kode & nama (baris 1-2) ---
    codes, labels, keep = [], {}, []
    for j in range(2, raw.shape[1]):
        col_letter = get_column_letter(j + 1)
        body_has = raw.iloc[2:, j].notna().any()
        code = raw.iat[1, j]
        if pd.isna(code) or str(code).strip() == "":
            if body_has:
                errors.append({"lokasi": f"Sel {col_letter}2", "pesan": "Kode indikator kosong padahal kolom berisi data. Isi kode singkat tanpa spasi, misalnya GPDRB."})
            continue
        code = str(code).strip()
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", code):
            errors.append({"lokasi": f"Sel {col_letter}2", "pesan": f"Kode \"{code}\" tidak valid. Pakai huruf, angka, atau garis bawah, diawali huruf, tanpa spasi."})
            continue
        if code in codes:
            errors.append({"lokasi": f"Sel {col_letter}2", "pesan": f"Kode \"{code}\" dipakai lebih dari sekali."})
            continue
        codes.append(code)
        keep.append(j)
        lab = raw.iat[0, j]
        labels[code] = code if pd.isna(lab) else str(lab).strip()
    if not codes:
        errors.append({"lokasi": "Baris 2", "pesan": "Tidak ada kode indikator di baris 2."})
        return res

    # --- periode (kolom A-B) ---
    body = raw.iloc[2:].copy()
    body = body[~body.iloc[:, [0, 1] + keep].isna().all(axis=1)]
    has_val = body.iloc[:, keep].notna().any(axis=1)
    years = pd.to_numeric(body[0], errors="coerce").ffill()
    periods, rows_ok = [], []
    for (ridx, row), yr in zip(body.iterrows(), years):
        excel_row = ridx + 1
        q = str(row[1]).strip().upper().replace(" ", "") if pd.notna(row[1]) else ""
        if pd.isna(yr):
            errors.append({"lokasi": f"Sel A{excel_row}", "pesan": "Tahun belum diisi. Isi tahun minimal di baris Q1."})
            continue
        if q not in Q_OK:
            errors.append({"lokasi": f"Sel B{excel_row}", "pesan": f"Triwulan tertulis \"{row[1]}\". Isi hanya Q1, Q2, Q3, atau Q4."})
            continue
        periods.append(pd.Period(f"{int(yr)}{q}", "Q"))
        rows_ok.append(ridx)
    if errors:
        return res
    for a, b, r in zip(periods, periods[1:], rows_ok[1:]):
        if b != a + 1:
            errors.append({"lokasi": f"Baris {r + 1}", "pesan": f"Periode tidak berurutan: setelah {a} seharusnya {a + 1}, tertulis {b}."})
            break

    # --- nilai ---
    data = {}
    for j, code in zip(keep, codes):
        col_letter = get_column_letter(j + 1)
        vals = []
        for ridx, p in zip(rows_ok, periods):
            v = raw.iat[ridx, j]
            if pd.isna(v) or (isinstance(v, str) and v.strip() == ""):
                vals.append(np.nan)
                continue
            try:
                vals.append(float(str(v).replace(",", ".")) if isinstance(v, str) else float(v))
            except ValueError:
                errors.append({"lokasi": f"Sel {col_letter}{ridx + 1} ({code}, {p})", "pesan": f"Berisi teks \"{v}\". Kosongkan sel atau isi angka."})
                vals.append(np.nan)
        data[code] = vals
    if errors:
        return res
    df = pd.DataFrame(data, index=pd.PeriodIndex(periods, freq="Q")).dropna(how="all")
    if df.empty:
        errors.append({"lokasi": "Data", "pesan": "Tidak ada angka yang terbaca. Isi data di sheet \"Data\" lalu unggah ulang."})
        return res

    # --- peringatan ---
    last_all = max(df[c].last_valid_index() for c in df if df[c].notna().any())
    for c in df:
        s = df[c]
        if s.notna().sum() == 0:
            warnings.append({"lokasi": c, "pesan": "Kolom kosong, tidak akan dipakai."})
            continue
        lv, fv = s.last_valid_index(), s.first_valid_index()
        if lv < last_all:
            warnings.append({"lokasi": c, "pesan": f"Data berhenti di {lv}, lebih awal dari periode terakhir ({last_all}). Nilai yang kosong akan diisi otomatis saat proyeksi."})
        gaps = s.loc[fv:lv].isna().sum()
        if gaps:
            warnings.append({"lokasi": c, "pesan": f"Ada {gaps} sel kosong di tengah data. Akan diisi interpolasi."})
        out = deteksi_outlier(s)
        if len(out):
            contoh = ", ".join(f"{p} ({v:.2f})" for p, v in list(out.items())[:3])
            warnings.append({"lokasi": c, "pesan": f"Kemungkinan ada outlier: {len(out)} nilai, misalnya {contoh}."})
    nmax = int(df.notna().sum().max())
    if nmax < MIN_OBS:
        errors.append({"lokasi": "Data", "pesan": f"Data terlalu pendek ({nmax} triwulan). Minimal {MIN_OBS} triwulan agar model bisa diuji."})
        return res
    if nmax < 30:
        warnings.append({"lokasi": "Data", "pesan": f"Data {nmax} triwulan. Jumlah uji backtest akan dikurangi otomatis."})

    res.update(ok=True, df=df, labels=labels, info={
        "periode_awal": str(df.index[0]), "periode_akhir": str(last_all), "n_obs": int(len(df.loc[:last_all])),
        "kolom": codes, "sel_kosong": int(df.loc[:last_all].isna().sum().sum())})
    return res


def tebak_pasangan_level(codes, df):
    """Tebak kolom level: nilai besar & selalu positif (mis. PDRB_ADHK)."""
    cand = []
    for c in codes:
        s = df[c].dropna()
        if len(s) and (s > 0).all() and s.median() > 1000:
            cand.append(c)
    return cand


def deteksi_outlier(s, batas=4.0):
    """Nilai dengan |z robust| > batas (median dan MAD). Mengembalikan Series periode -> nilai."""
    x = s.dropna()
    if len(x) < 8:
        return x.iloc[0:0]
    med = x.median()
    mad = 1.4826 * (x - med).abs().median()
    if not mad > 0:
        return x.iloc[0:0]
    z = (x - med) / mad
    return x[z.abs() > batas]
